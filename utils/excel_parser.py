"""
Excel (.xlsx) Question Bank Parser and Exporter (Phase 17).
Supports:
- Parsing and row-by-row validation of .xlsx question banks.
- Flexible column header resolution.
- Reading passage data (passage_title, passage_text).
- Listening audio references (audio_file, transcript, duration).
- Rule-based difficulty classification for imported items.
- Formatting and exporting filtered question banks to .xlsx.
- Template generation with sample rows.
"""

from typing import Optional, Any, BinaryIO
from pathlib import Path
import io
import re
import pandas as pd
import openpyxl

from utils.difficulty import classify_difficulty


STANDARD_COLUMNS = [
    "category",
    "skill_number",
    "skill_name",
    "question_order",
    "topic",
    "subtopic",
    "question_type",
    "question_text",
    "option_a",
    "option_b",
    "option_c",
    "option_d",
    "correct_answer",
    "explanation",
    "difficulty",
    "status",
    "passage_title",
    "passage_text",
    "audio_file",
    "transcript",
    "duration",
]

# Aliases for flexible header matching
COLUMN_ALIASES = {
    "category": ["category", "kategori", "cat"],
    "skill_number": ["skill_number", "skill_no", "skill", "nomor_skill", "no_skill", "nomor skill"],
    "skill_name": ["skill_name", "nama_skill", "skill_title", "kemampuan", "nama skill"],
    "question_order": ["question_order", "order", "urutan", "no", "number"],
    "topic": ["topic", "topik", "subject"],
    "subtopic": ["subtopic", "sub_topic", "sub-topic"],
    "question_type": ["question_type", "type", "tipe", "q_type"],
    "question_text": ["question_text", "question", "prompt", "pertanyaan", "soal"],
    "option_a": ["option_a", "option a", "opt_a", "a", "choice_a"],
    "option_b": ["option_b", "option b", "opt_b", "b", "choice_b"],
    "option_c": ["option_c", "option c", "opt_c", "c", "choice_c"],
    "option_d": ["option_d", "option d", "opt_d", "d", "choice_d"],
    "correct_answer": ["correct_answer", "correct answer", "answer", "key", "kunci", "jawaban"],
    "explanation": ["explanation", "pembahasan", "explain", "keterangan"],
    "difficulty": ["difficulty", "level", "tingkat"],
    "status": ["status"],
    "passage_title": ["passage_title", "title", "judul_bacaan"],
    "passage_text": ["passage_text", "passage", "text", "bacaan", "teks"],
    "audio_file": ["audio_file", "audio", "file_audio", "audio_path"],
    "transcript": ["transcript", "transkrip"],
    "duration": ["duration", "durasi"],
}


def normalize_column_name(col_name: str) -> Optional[str]:
    """Resolve raw header string to standard column name using aliases."""
    clean = re.sub(r"[_\s\-]+", " ", str(col_name).strip().lower())
    for std_name, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if clean == alias or clean == alias.replace("_", " "):
                return std_name
    return None


def parse_excel_content(
    file_source: bytes | BinaryIO | str | Path,
    default_category: str = "Grammar",
) -> dict[str, Any]:
    """
    Parse an Excel (.xlsx) question bank workbook.
    Validates every question item and attaches difficulty heuristics.
    Returns:
    {
        "total_detected": int,
        "valid_count": int,
        "invalid_count": int,
        "global_topic": str,
        "questions": list[dict],
    }
    """
    # Read into pandas DataFrame
    if isinstance(file_source, bytes):
        stream = io.BytesIO(file_source)
        df = pd.read_excel(stream, engine="openpyxl")
    else:
        df = pd.read_excel(file_source, engine="openpyxl")

    # Rename columns to standard internal names
    rename_map = {}
    for col in df.columns:
        norm = normalize_column_name(str(col))
        if norm:
            rename_map[col] = norm

    df = df.rename(columns=rename_map)

    questions_raw: list[dict[str, Any]] = []
    global_topics: set[str] = set()

    for idx, row in df.iterrows():
        # Skip rows that are completely empty
        if row.isna().all():
            continue

        q_num = idx + 1
        errors: list[str] = []

        # 1. Category
        raw_cat = str(row.get("category", "")).strip() if pd.notna(row.get("category")) else ""
        if raw_cat:
            cat_norm = raw_cat.capitalize()
            if cat_norm not in ["Grammar", "Reading", "Listening"]:
                errors.append(f"Invalid category '{raw_cat}'. Must be Grammar, Reading, or Listening.")
                cat_clean = default_category
            else:
                cat_clean = cat_norm
        else:
            cat_clean = default_category

        # 2. Question Text
        raw_text = str(row.get("question_text", "")).strip() if pd.notna(row.get("question_text")) else ""
        if not raw_text:
            if cat_clean == "Listening":
                raw_text = "Listen to the audio recording and choose the correct answer."
            else:
                errors.append("Question prompt text is missing or empty.")

        # 3. Options A-D
        opt_a = str(row.get("option_a", "")).strip() if pd.notna(row.get("option_a")) else ""
        opt_b = str(row.get("option_b", "")).strip() if pd.notna(row.get("option_b")) else ""
        opt_c = str(row.get("option_c", "")).strip() if pd.notna(row.get("option_c")) else ""
        opt_d = str(row.get("option_d", "")).strip() if pd.notna(row.get("option_d")) else ""

        missing_opts = []
        if not opt_a: missing_opts.append("A")
        if not opt_b: missing_opts.append("B")
        if not opt_c: missing_opts.append("C")
        if not opt_d: missing_opts.append("D")

        if missing_opts:
            errors.append(f"Missing option(s): {', '.join(missing_opts)}")
        elif len({opt_a.lower(), opt_b.lower(), opt_c.lower(), opt_d.lower()}) < 4:
            errors.append("Duplicate choices detected among options A, B, C, D.")

        # 4. Correct Answer
        raw_ans = str(row.get("correct_answer", "")).strip().upper() if pd.notna(row.get("correct_answer")) else ""
        if not raw_ans:
            errors.append("Answer key is missing.")
        elif raw_ans not in ["A", "B", "C", "D"]:
            errors.append(f"Invalid answer key '{raw_ans}'. Must be A, B, C, or D.")

        # 5. Difficulty
        raw_diff = str(row.get("difficulty", "")).strip().capitalize() if pd.notna(row.get("difficulty")) else ""
        if raw_diff in ["Easy", "Medium", "Hard"]:
            assigned_diff = raw_diff
        else:
            assigned_diff = None  # will derive from heuristic

        # 6. Status
        raw_status = str(row.get("status", "")).strip().capitalize() if pd.notna(row.get("status")) else ""
        if raw_status in ["Draft", "Active", "Inactive", "Archived"]:
            q_status = raw_status
        else:
            q_status = "Draft"

        # 7. Topic & Subtopic
        topic = str(row.get("topic", "")).strip() if pd.notna(row.get("topic")) else ""
        subtopic = str(row.get("subtopic", "")).strip() if pd.notna(row.get("subtopic")) else (topic or "General")
        if topic:
            global_topics.add(topic)

        # 8. Reading Passage
        passage_title = str(row.get("passage_title", "")).strip() if pd.notna(row.get("passage_title")) else "Imported Passage"
        passage_text = str(row.get("passage_text", "")).strip() if pd.notna(row.get("passage_text")) else ""

        if cat_clean == "Reading" and not passage_text:
            errors.append("Reading questions require passage text ('passage_text' column is empty).")

        # 9. Listening Audio
        audio_file = str(row.get("audio_file", "")).strip() if pd.notna(row.get("audio_file")) else ""
        transcript = str(row.get("transcript", "")).strip() if pd.notna(row.get("transcript")) else ""
        raw_dur = row.get("duration")
        duration = float(raw_dur) if pd.notna(raw_dur) and isinstance(raw_dur, (int, float)) else None

        # 10. Explanation
        explanation = str(row.get("explanation", "")).strip() if pd.notna(row.get("explanation")) else "No explanation provided."

        # 11. Skill and Ordering
        raw_snum = row.get("skill_number")
        skill_number = 1
        if pd.notna(raw_snum):
            try:
                skill_number = int(raw_snum)
            except (ValueError, TypeError):
                skill_number = 1

        raw_sname = row.get("skill_name")
        skill_name = str(raw_sname).strip() if pd.notna(raw_sname) and str(raw_sname).strip() else (topic or f"Skill {skill_number}")

        raw_qorder = row.get("question_order")
        question_order = q_num
        if pd.notna(raw_qorder):
            try:
                question_order = int(raw_qorder)
            except (ValueError, TypeError):
                question_order = q_num

        # Rule-based difficulty heuristic evaluation (Phase 13)
        diff_res = classify_difficulty(
            category=cat_clean,
            question_text=raw_text,
            options=[opt_a, opt_b, opt_c, opt_d],
            passage_text=passage_text,
            duration=duration,
            transcript=transcript,
            explanation=explanation,
        )

        final_difficulty = assigned_diff or diff_res.suggested_difficulty

        is_valid = (len(errors) == 0)

        questions_raw.append({
            "question_number": q_num,
            "category": cat_clean,
            "skill_number": skill_number,
            "skill_name": skill_name,
            "question_order": question_order,
            "topic": topic or subtopic,
            "subtopic": subtopic,
            "question_type": str(row.get("question_type", "Multiple Choice")),
            "question_text": raw_text,
            "option_a": opt_a,
            "option_b": opt_b,
            "option_c": opt_c,
            "option_d": opt_d,
            "correct_answer": raw_ans,
            "explanation": explanation,
            "difficulty": final_difficulty,
            "suggested_difficulty": diff_res.suggested_difficulty,
            "difficulty_score": diff_res.difficulty_score,
            "reason": diff_res.reason,
            "difficulty_reason": diff_res.reason,
            "final_difficulty": final_difficulty,
            "status": q_status,
            "passage_title": passage_title,
            "passage_text": passage_text,
            "audio_file": audio_file,
            "transcript": transcript,
            "duration": duration,
            "is_valid": is_valid,
            "validation_errors": errors,
        })

    valid_qs = [q for q in questions_raw if q["is_valid"]]
    invalid_qs = [q for q in questions_raw if not q["is_valid"]]

    return {
        "total_detected": len(questions_raw),
        "valid_count": len(valid_qs),
        "invalid_count": len(invalid_qs),
        "global_topic": ", ".join(list(global_topics)[:3]) if global_topics else "General",
        "questions": questions_raw,
    }


SKILL_TEMPLATE_COLUMNS = [
    "category",
    "skill_number",
    "skill_name",
    "question_text",
    "option_a",
    "option_b",
    "option_c",
    "option_d",
    "correct_answer",
    "explanation",
    "passage_title",
    "passage_text",
    "audio_file",
    "transcript",
]


def export_questions_to_excel(
    questions: list[dict[str, Any]],
    columns: Optional[list[str]] = None,
) -> bytes:
    """
    Format and export question records to an Excel (.xlsx) spreadsheet binary.
    """
    target_cols = columns or STANDARD_COLUMNS
    export_rows = []
    for q in questions:
        row = {}
        for col in target_cols:
            if col == "category":
                row["category"] = q.get("category_name") or q.get("category", "Grammar")
            elif col == "skill_number":
                row["skill_number"] = q.get("skill_number", 1)
            elif col == "skill_name":
                row["skill_name"] = q.get("skill_name", f"Skill {q.get('skill_number', 1)}")
            elif col == "correct_answer":
                row["correct_answer"] = q.get("correct_answer", "A")
            elif col == "explanation":
                row["explanation"] = q.get("explanation", "")
            else:
                row[col] = q.get(col, "")
        export_rows.append(row)

    df = pd.DataFrame(export_rows, columns=target_cols)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Question Bank")

        # Auto-adjust column widths for readability
        worksheet = writer.sheets["Question Bank"]
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = min(45, max(max_len + 3, 14))

    return output.getvalue()


def generate_excel_import_template() -> bytes:
    """Generate a clean sample Excel template for skill-based question bank import (without difficulty column)."""
    sample_data = [
        {
            "category": "Grammar",
            "skill_number": 1,
            "skill_name": "Subject-Verb Agreement",
            "question_text": "The list of new committee members ___ posted on the notice board.",
            "option_a": "are",
            "option_b": "is",
            "option_c": "were",
            "option_d": "be",
            "correct_answer": "B",
            "explanation": "Subjek kalimat adalah kata benda tunggal 'list' (bukan 'members'), sehingga membutuhkan kata kerja singular 'is'.",
            "passage_title": "",
            "passage_text": "",
            "audio_file": "",
            "transcript": "",
        },
        {
            "category": "Reading",
            "skill_number": 1,
            "skill_name": "Main Ideas and Details",
            "question_text": "According to paragraph 1, what enables deep sea organisms to produce light?",
            "option_a": "Specialized photophores",
            "option_b": "Solar radiation absorption",
            "option_c": "Ambient water temperature",
            "option_d": "Hydrostatic pressure",
            "correct_answer": "A",
            "explanation": "The text states bioluminescence is generated by light-emitting photophores.",
            "passage_title": "Abyssal Marine Life",
            "passage_text": "In deep oceanic trenches, organisms exhibit bioluminescence produced by specialized light-emitting photophores.",
            "audio_file": "",
            "transcript": "",
        },
        {
            "category": "Listening",
            "skill_number": 1,
            "skill_name": "Short Dialogues & Details",
            "question_text": "What time will the central library open on Saturday morning?",
            "option_a": "At 7:00 AM",
            "option_b": "At 8:30 AM",
            "option_c": "At 9:00 AM",
            "option_d": "At 10:00 AM",
            "correct_answer": "C",
            "explanation": "The audio announcement specifically announces Saturday opening hours at 9:00 AM.",
            "passage_title": "",
            "passage_text": "",
            "audio_file": "audio/listening/library_hours.mp3",
            "transcript": "Attention students, the central library will open at 9:00 AM this Saturday.",
        },
    ]

    return export_questions_to_excel(sample_data, columns=SKILL_TEMPLATE_COLUMNS)


def export_users_to_excel(users: list[dict[str, Any]]) -> bytes:
    """
    Format and export participant user directory records to an Excel (.xlsx) spreadsheet.
    """
    rows = []
    for u in users:
        avg_display = f"{u['average_score']}%" if u.get("average_score") is not None else "No scores yet"
        created_display = u["created_at"][:10] if u.get("created_at") else "-"
        rows.append({
            "User ID": u.get("id"),
            "Full Name": u.get("name", ""),
            "Email Address": u.get("email", ""),
            "Institution": u.get("institution") or "-",
            "Program / Role": u.get("program") or "-",
            "Registration Date": created_display,
            "Total Tests Completed": u.get("total_attempts", 0),
            "Average Score": avg_display,
        })

    cols = [
        "User ID",
        "Full Name",
        "Email Address",
        "Institution",
        "Program / Role",
        "Registration Date",
        "Total Tests Completed",
        "Average Score",
    ]
    df = pd.DataFrame(rows, columns=cols)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Participant Directory")

        worksheet = writer.sheets["Participant Directory"]
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = min(45, max(max_len + 3, 14))

    return output.getvalue()


def export_performance_highlights_to_excel(highlights_dict: dict[str, list[dict[str, Any]]]) -> bytes:
    """
    Export question performance highlights across multiple tabs (Most Frequent, Highest Accuracy,
    Lowest Accuracy, Attention Needed) to a single multi-sheet Excel (.xlsx) file.
    """
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        cols = [
            "ID",
            "Category",
            "Skill",
            "Urutan",
            "Question Text",
            "Attempts",
            "Correct",
            "Incorrect",
            "Accuracy",
            "Observation",
        ]

        has_written_sheet = False
        for sheet_name, items in highlights_dict.items():
            safe_sheet_name = re.sub(r"[\\/*?:\[\]]", "", sheet_name)[:31]
            rows = []
            for q in items:
                skill_str = f"Skill {q.get('skill_number', 1)}: {q.get('skill_name', '')}"
                rows.append({
                    "ID": q.get("id"),
                    "Category": q.get("category", ""),
                    "Skill": skill_str,
                    "Urutan": f"#{q.get('question_order', 1)}",
                    "Question Text": q.get("question_text", ""),
                    "Attempts": q.get("attempts", 0),
                    "Correct": q.get("correct_count", 0),
                    "Incorrect": q.get("incorrect_count", 0),
                    "Accuracy": f"{q.get('accuracy', 0)}%",
                    "Observation": q.get("admin_note", ""),
                })

            df = pd.DataFrame(rows, columns=cols)
            df.to_excel(writer, index=False, sheet_name=safe_sheet_name)
            has_written_sheet = True

            worksheet = writer.sheets[safe_sheet_name]
            for col in worksheet.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                worksheet.column_dimensions[col_letter].width = min(50, max(max_len + 3, 12))

        if not has_written_sheet:
            pd.DataFrame(columns=cols).to_excel(writer, index=False, sheet_name="Highlights")

    return output.getvalue()


def export_question_performance_to_excel(analytics_items: list[dict[str, Any]]) -> bytes:
    """
    Format and export complete question performance table to an Excel (.xlsx) spreadsheet.
    """
    cols = [
        "ID",
        "Category",
        "Skill Number",
        "Skill Name",
        "Urutan",
        "Question Prompt",
        "Total Attempts",
        "Correct Count",
        "Incorrect Count",
        "Observed Accuracy",
        "Needs Review",
        "Admin Observation",
    ]

    rows = []
    for q in analytics_items:
        rows.append({
            "ID": q.get("id"),
            "Category": q.get("category", ""),
            "Skill Number": q.get("skill_number", 1),
            "Skill Name": q.get("skill_name", ""),
            "Urutan": f"#{q.get('question_order', 1)}",
            "Question Prompt": q.get("question_text", ""),
            "Total Attempts": q.get("attempts", 0),
            "Correct Count": q.get("correct_count", 0),
            "Incorrect Count": q.get("incorrect_count", 0),
            "Observed Accuracy": f"{q.get('accuracy', 0)}%",
            "Needs Review": "Yes (Review Recommended)" if q.get("needs_review") else "No",
            "Admin Observation": q.get("admin_note", ""),
        })

    df = pd.DataFrame(rows, columns=cols)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Question Performance")

        worksheet = writer.sheets["Question Performance"]
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = min(55, max(max_len + 3, 12))

    return output.getvalue()


def export_assessment_results_to_excel(results: list[dict[str, Any]]) -> bytes:
    """
    Format and export historical assessment attempt records to an Excel (.xlsx) spreadsheet.
    """
    cols = [
        "Attempt ID",
        "Participant Name",
        "Email Address",
        "Institution",
        "Category",
        "Skill",
        "Score",
        "Correct Answers",
        "Incorrect Answers",
        "Duration",
        "Started At",
        "Completed At",
    ]

    rows = []
    for r in results:
        dur = r.get("duration_seconds")
        dur_str = f"{dur // 60}m {dur % 60}s" if dur else "N/A"
        started_str = r["started_at"][:19] if r.get("started_at") else "-"
        completed_str = r["completed_at"][:19] if r.get("completed_at") else "In Progress"

        rows.append({
            "Attempt ID": r.get("id"),
            "Participant Name": r.get("user_name", ""),
            "Email Address": r.get("user_email", ""),
            "Institution": r.get("institution") or "-",
            "Category": r.get("category", ""),
            "Skill": r.get("difficulty", ""),
            "Score": f"{r.get('score', 0)}%",
            "Correct Answers": f"{r.get('correct_answers', 0)} / {r.get('total_questions', 0)}",
            "Incorrect Answers": r.get("incorrect_answers", 0),
            "Duration": dur_str,
            "Started At": started_str,
            "Completed At": completed_str,
        })

    df = pd.DataFrame(rows, columns=cols)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Assessment Results")

        worksheet = writer.sheets["Assessment Results"]
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = min(45, max(max_len + 3, 12))

    return output.getvalue()
