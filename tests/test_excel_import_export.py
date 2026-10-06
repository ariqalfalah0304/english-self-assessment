"""
Tests for Phase 17: Excel (.xlsx) Import and Export.
Validates:
- Excel parsing with alias tolerance and multi-format column handling
- Row validation: missing fields, duplicates, invalid answer keys
- Reading passage requirements and Listening audio file linking
- Preview data structure and rule-based difficulty suggestion
- Export of filtered questions to formatted .xlsx
- Template generation
- Database batch import integrity, passage creation, audio creation, and history logging
"""

import io
import unittest
import tempfile
import shutil
from pathlib import Path
import pandas as pd
import openpyxl

from database.db import init_db, get_connection
from database import queries
from database.models import Question
from utils.excel_parser import (
    parse_excel_content,
    export_questions_to_excel,
    export_users_to_excel,
    export_performance_highlights_to_excel,
    export_question_performance_to_excel,
    export_assessment_results_to_excel,
    generate_excel_import_template,
    STANDARD_COLUMNS,
)
from services.import_service import execute_import_batch, fetch_import_history_logs


class TestExcelImportExport(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.db_path = Path(self.temp_dir) / "test_assessment.db"
        init_db(self.db_path)

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_generate_and_parse_template(self):
        """Verify Excel import template generation and structure."""
        template_bytes = generate_excel_import_template()
        self.assertIsInstance(template_bytes, bytes)
        self.assertGreater(len(template_bytes), 1000)

        # Parse the template back
        parsed = parse_excel_content(template_bytes)
        self.assertEqual(parsed["total_detected"], 3)
        self.assertEqual(parsed["valid_count"], 3)
        self.assertEqual(parsed["invalid_count"], 0)

        categories = [q["category"] for q in parsed["questions"]]
        self.assertIn("Grammar", categories)
        self.assertIn("Reading", categories)
        self.assertIn("Listening", categories)

        # Verify Reading has passage
        reading_q = next(q for q in parsed["questions"] if q["category"] == "Reading")
        self.assertTrue(len(reading_q["passage_text"]) > 10)

        # Verify Listening has audio file
        listening_q = next(q for q in parsed["questions"] if q["category"] == "Listening")
        self.assertTrue("audio" in listening_q["audio_file"])

    def test_parse_valid_custom_excel(self):
        """Verify parsing of valid rows across categories with difficulty heuristic."""
        rows = [
            {
                "category": "Grammar",
                "topic": "Tenses",
                "subtopic": "Present Perfect",
                "question_type": "Multiple Choice",
                "question_text": "She has ___ to Paris three times this year.",
                "option_a": "go",
                "option_b": "went",
                "option_c": "gone",
                "option_d": "going",
                "correct_answer": "C",
                "explanation": "Present perfect requires past participle.",
                "difficulty": "Easy",
                "status": "Draft",
            },
            {
                "category": "Reading",
                "topic": "Science",
                "subtopic": "Geology",
                "question_type": "Main Idea",
                "question_text": "What is the primary factor leading to tectonic shift?",
                "option_a": "Mantle convection currents",
                "option_b": "Atmospheric pressure",
                "option_c": "Oceanic salinity",
                "option_d": "Gravitational tides",
                "correct_answer": "A",
                "explanation": "Convection in the mantle drives plate tectonics.",
                "difficulty": "Medium",
                "status": "Draft",
                "passage_title": "Plate Tectonics",
                "passage_text": "Thermal convection in the Earth's mantle drives the slow movement of lithospheric plates.",
            },
            {
                "category": "Listening",
                "topic": "University",
                "subtopic": "Lecture",
                "question_type": "Detail",
                "question_text": "When is the lab report due according to the professor?",
                "option_a": "Monday noon",
                "option_b": "Tuesday 5 PM",
                "option_c": "Wednesday 9 AM",
                "option_d": "Friday midnight",
                "correct_answer": "B",
                "explanation": "The lecture specifies Tuesday 5 PM.",
                "difficulty": "Hard",
                "status": "Draft",
                "audio_file": "audio/listening/lab_announcement.mp3",
                "transcript": "Please submit your completed lab reports by Tuesday at 5 PM sharp.",
                "duration": 22.0,
            }
        ]

        df = pd.DataFrame(rows)
        stream = io.BytesIO()
        df.to_excel(stream, index=False, engine="openpyxl")
        excel_bytes = stream.getvalue()

        parsed = parse_excel_content(excel_bytes)
        self.assertEqual(parsed["total_detected"], 3)
        self.assertEqual(parsed["valid_count"], 3)
        self.assertEqual(parsed["invalid_count"], 0)

        for q in parsed["questions"]:
            self.assertTrue(q["is_valid"])
            self.assertIn(q["suggested_difficulty"], ["Easy", "Medium", "Hard"])
            self.assertIsNotNone(q["difficulty_score"])
            self.assertTrue(len(q["reason"]) > 0)

    def test_validation_rejects_invalid_excel_rows(self):
        """Verify strict row validation for missing fields, duplicates, and bad keys."""
        rows = [
            # 1. Missing question text
            {
                "category": "Grammar",
                "question_text": "",
                "option_a": "A", "option_b": "B", "option_c": "C", "option_d": "D",
                "correct_answer": "A",
            },
            # 2. Missing option D
            {
                "category": "Grammar",
                "question_text": "Valid question prompt?",
                "option_a": "Apple", "option_b": "Banana", "option_c": "Cherry", "option_d": "",
                "correct_answer": "A",
            },
            # 3. Duplicate options
            {
                "category": "Grammar",
                "question_text": "Which is an animal?",
                "option_a": "Cat", "option_b": "Dog", "option_c": "Cat", "option_d": "Horse",
                "correct_answer": "A",
            },
            # 4. Invalid correct answer key
            {
                "category": "Grammar",
                "question_text": "Select option E?",
                "option_a": "A", "option_b": "B", "option_c": "C", "option_d": "D",
                "correct_answer": "E",
            },
            # 5. Reading question missing passage
            {
                "category": "Reading",
                "question_text": "What does the passage say?",
                "option_a": "A", "option_b": "B", "option_c": "C", "option_d": "D",
                "correct_answer": "B",
                "passage_text": "",
            },
            # 6. Listening question missing answer key
            {
                "category": "Listening",
                "question_text": "What did the speaker say?",
                "option_a": "A", "option_b": "B", "option_c": "C", "option_d": "D",
                "correct_answer": "",
                "audio_file": "",
            },
            # 7. Invalid category
            {
                "category": "Calculus",
                "question_text": "What is the derivative of x^2?",
                "option_a": "2x", "option_b": "x", "option_c": "x^3", "option_d": "2",
                "correct_answer": "A",
            },
        ]

        df = pd.DataFrame(rows)
        stream = io.BytesIO()
        df.to_excel(stream, index=False, engine="openpyxl")

        parsed = parse_excel_content(stream.getvalue())
        self.assertEqual(parsed["total_detected"], 7)
        self.assertEqual(parsed["valid_count"], 0)
        self.assertEqual(parsed["invalid_count"], 7)

        # Check specific validation errors
        errs_combined = " ".join([" ".join(q["validation_errors"]) for q in parsed["questions"]])
        self.assertIn("Question prompt text is missing", errs_combined)
        self.assertIn("Missing option(s)", errs_combined)
        self.assertIn("Duplicate choices detected", errs_combined)
        self.assertIn("Invalid answer key", errs_combined)
        self.assertIn("Reading questions require passage text", errs_combined)
        self.assertIn("Invalid category", errs_combined)

    def test_export_filtered_questions_to_excel(self):
        """Verify export of question records produces valid .xlsx spreadsheet."""
        mock_questions = [
            {
                "id": 101,
                "category_name": "Grammar",
                "subtopic": "Passive Voice",
                "question_type": "Multiple Choice",
                "question_text": "The bridge was designed by Eiffel.",
                "option_a": "was designed",
                "option_b": "is designing",
                "option_c": "were designed",
                "option_d": "designed",
                "correct_answer": "A",
                "explanation": "Past passive voice.",
                "difficulty": "Medium",
                "status": "Active",
            },
            {
                "id": 102,
                "category_name": "Reading",
                "subtopic": "History",
                "question_type": "Inference",
                "question_text": "Why did the construction halt?",
                "option_a": "Funding shortage",
                "option_b": "Severe weather",
                "option_c": "Labor strike",
                "option_d": "Permit denial",
                "correct_answer": "B",
                "explanation": "Winter blizzards stopped the builders.",
                "difficulty": "Hard",
                "status": "Draft",
                "passage_title": "The Northern Bridge",
                "passage_text": "Construction was interrupted by heavy snowfalls during the winter of 1888.",
            },
        ]

        excel_bytes = export_questions_to_excel(mock_questions)
        self.assertIsInstance(excel_bytes, bytes)

        # Read back with pandas
        df_read = pd.read_excel(io.BytesIO(excel_bytes), engine="openpyxl")
        self.assertEqual(len(df_read), 2)
        self.assertListEqual(list(df_read.columns), STANDARD_COLUMNS)
        self.assertEqual(df_read.iloc[0]["category"], "Grammar")
        self.assertEqual(df_read.iloc[1]["passage_title"], "The Northern Bridge")
        self.assertEqual(df_read.iloc[1]["correct_answer"], "B")

    def test_batch_import_database_integrity(self):
        """Verify executing an Excel import batch commits to database with foreign keys & drafts."""
        valid_items = [
            {
                "is_valid": True,
                "category": "Grammar",
                "topic": "Prepositions",
                "subtopic": "Time",
                "question_type": "Multiple Choice",
                "question_text": "We will meet ___ noon tomorrow.",
                "option_a": "at", "option_b": "in", "option_c": "on", "option_d": "by",
                "correct_answer": "A",
                "explanation": "Use 'at' with noon.",
                "final_difficulty": "Easy",
                "status": "Draft",
            },
            {
                "is_valid": True,
                "category": "Reading",
                "topic": "Ecology",
                "subtopic": "Forestry",
                "question_type": "Vocabulary",
                "question_text": "The word 'canopy' refers to which part?",
                "option_a": "The roots", "option_b": "The upper layer", "option_c": "The forest floor", "option_d": "The soil",
                "correct_answer": "B",
                "explanation": "Canopy is the uppermost level of foliage.",
                "final_difficulty": "Medium",
                "status": "Draft",
                "passage_title": "Rainforest Canopy",
                "passage_text": "The forest canopy houses diverse arboreal species living high above ground.",
            },
            {
                "is_valid": True,
                "category": "Listening",
                "topic": "Campus",
                "subtopic": "Office Hours",
                "question_type": "Detail",
                "question_text": "Where is Professor Smith's office located?",
                "option_a": "Room 101", "option_b": "Hall B", "option_c": "Room 405", "option_d": "Tower C",
                "correct_answer": "C",
                "explanation": "The conversation states Room 405.",
                "final_difficulty": "Hard",
                "status": "Draft",
                "audio_file": "audio/listening/smith_office.mp3",
                "transcript": "You can find me in Room 405 on Wednesdays.",
                "duration": 15.5,
            },
            # Invalid item that should be skipped
            {
                "is_valid": False,
                "question_text": "Broken item without options",
                "validation_errors": ["Missing options"],
            }
        ]

        detected, imported, failed = execute_import_batch(
            file_name="excel_test_upload.xlsx",
            questions=valid_items,
            category_name="Grammar",
            default_status="Draft",
            imported_by="test_admin",
            db_path=self.db_path,
        )

        self.assertEqual(detected, 4)
        self.assertEqual(imported, 3)
        self.assertEqual(failed, 1)

        # Verify DB records
        with get_connection(self.db_path) as conn:
            # 1. Grammar question
            grammar_cat = queries.get_category_by_name(conn, "Grammar")
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM questions WHERE category_id = ? AND status = 'Draft';", (grammar_cat.id,))
            g_rows = cursor.fetchall()
            self.assertEqual(len(g_rows), 1)
            self.assertEqual(g_rows[0][6], "We will meet ___ noon tomorrow.")

            # 2. Reading question & passage
            reading_cat = queries.get_category_by_name(conn, "Reading")
            cursor.execute("SELECT * FROM questions WHERE category_id = ?;", (reading_cat.id,))
            r_rows = cursor.fetchall()
            self.assertEqual(len(r_rows), 1)
            r_passage_id = r_rows[0][3]  # passage_id
            self.assertIsNotNone(r_passage_id)
            cursor.execute("SELECT title, passage_text FROM passages WHERE id = ?;", (r_passage_id,))
            pass_row = cursor.fetchone()
            self.assertEqual(pass_row[0], "Rainforest Canopy")

            # 3. Listening question & audio_files
            listening_cat = queries.get_category_by_name(conn, "Listening")
            cursor.execute("SELECT id FROM questions WHERE category_id = ?;", (listening_cat.id,))
            l_rows = cursor.fetchall()
            self.assertEqual(len(l_rows), 1)
            l_qid = l_rows[0][0]

            audio_rec = queries.get_audio_file_by_question_id(conn, l_qid)
            self.assertIsNotNone(audio_rec)
            self.assertEqual(audio_rec.file_path, "audio/listening/smith_office.mp3")
            self.assertEqual(audio_rec.transcript, "You can find me in Room 405 on Wednesdays.")
            self.assertEqual(audio_rec.duration, 15.5)

        # 4. Check import history
        logs = fetch_import_history_logs(limit=5, db_path=self.db_path)
        self.assertEqual(len(logs), 1)
        self.assertEqual(logs[0].file_name, "excel_test_upload.xlsx")
        self.assertEqual(logs[0].total_imported, 3)
        self.assertEqual(logs[0].total_failed, 1)

    def test_export_users_to_excel(self):
        """Verify participant user directory exports correctly to .xlsx."""
        mock_users = [
            {
                "id": 6,
                "name": "Ariq",
                "email": "ariq@gmail.com",
                "institution": "Universitas Nurdin Hamzah",
                "program": "Sistem Informasi",
                "created_at": "2026-10-02 10:00:00",
                "total_attempts": 10,
                "average_score": 52.5,
            },
            {
                "id": 7,
                "name": "Riko",
                "email": "riko@gmail.com",
                "institution": "Universitas Jambi",
                "program": "Pertanian",
                "created_at": "2026-10-03 14:00:00",
                "total_attempts": 0,
                "average_score": None,
            },
        ]
        excel_bytes = export_users_to_excel(mock_users)
        self.assertIsInstance(excel_bytes, bytes)

        df = pd.read_excel(io.BytesIO(excel_bytes), engine="openpyxl", sheet_name="Participant Directory")
        self.assertEqual(len(df), 2)
        self.assertEqual(df.iloc[0]["Full Name"], "Ariq")
        self.assertEqual(df.iloc[0]["Total Tests Completed"], 10)
        self.assertEqual(df.iloc[0]["Average Score"], "52.5%")
        self.assertEqual(df.iloc[1]["Full Name"], "Riko")
        self.assertEqual(df.iloc[1]["Average Score"], "No scores yet")

    def test_export_performance_highlights_to_excel(self):
        """Verify performance highlights multi-sheet export to .xlsx."""
        mock_item = {
            "id": 301,
            "category": "Grammar",
            "skill_number": 2,
            "skill_name": "Object of Preposition",
            "question_order": 3,
            "question_text": "The books on the table were expensive.",
            "attempts": 5,
            "correct_count": 4,
            "incorrect_count": 1,
            "accuracy": 80.0,
            "admin_note": "Accuracy: 80.0% from 5 attempt(s).",
        }
        highlights = {
            "Frequently Attempted": [mock_item],
            "Highest Accuracy": [mock_item],
            "Lowest Accuracy": [],
            "Attention Needed": [],
        }
        excel_bytes = export_performance_highlights_to_excel(highlights)
        self.assertIsInstance(excel_bytes, bytes)

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertIn("Frequently Attempted", wb.sheetnames)
        self.assertIn("Highest Accuracy", wb.sheetnames)

        df_freq = pd.read_excel(io.BytesIO(excel_bytes), engine="openpyxl", sheet_name="Frequently Attempted")
        self.assertEqual(len(df_freq), 1)
        self.assertEqual(df_freq.iloc[0]["ID"], 301)
        self.assertEqual(df_freq.iloc[0]["Accuracy"], "80.0%")

    def test_export_question_performance_to_excel(self):
        """Verify comprehensive question performance analytics exports to .xlsx."""
        mock_items = [
            {
                "id": 252,
                "category": "Grammar",
                "skill_number": 1,
                "skill_name": "Identifying Subjects and Verbs",
                "question_order": 4,
                "question_text": "The books on the table ___ mine.",
                "attempts": 2,
                "correct_count": 0,
                "incorrect_count": 2,
                "accuracy": 0.0,
                "needs_review": True,
                "admin_note": "Accuracy: 0.0% from 2 attempt(s).",
            }
        ]
        excel_bytes = export_question_performance_to_excel(mock_items)
        self.assertIsInstance(excel_bytes, bytes)

        df = pd.read_excel(io.BytesIO(excel_bytes), engine="openpyxl", sheet_name="Question Performance")
        self.assertEqual(len(df), 1)
        self.assertEqual(df.iloc[0]["ID"], 252)
        self.assertEqual(df.iloc[0]["Observed Accuracy"], "0.0%")
        self.assertEqual(df.iloc[0]["Needs Review"], "Yes (Review Recommended)")

    def test_export_assessment_results_to_excel(self):
        """Verify historical assessment attempt records export to .xlsx."""
        mock_results = [
            {
                "id": 22,
                "user_name": "Ariq",
                "user_email": "ariq@gmail.com",
                "institution": "Universitas Nurdin Hamzah",
                "category": "Grammar",
                "difficulty": "Ujian Gabungan",
                "score": 0.0,
                "correct_answers": 0,
                "total_questions": 50,
                "incorrect_answers": 0,
                "duration_seconds": None,
                "started_at": "2026-10-04 16:25:50",
                "completed_at": None,
            },
            {
                "id": 19,
                "user_name": "Ariq",
                "user_email": "ariq@gmail.com",
                "institution": "Universitas Nurdin Hamzah",
                "category": "Grammar",
                "difficulty": "Skill 3",
                "score": 100.0,
                "correct_answers": 4,
                "total_questions": 4,
                "incorrect_answers": 0,
                "duration_seconds": 17,
                "started_at": "2026-10-03 06:46:07",
                "completed_at": "2026-10-03 06:46:24",
            },
        ]
        excel_bytes = export_assessment_results_to_excel(mock_results)
        self.assertIsInstance(excel_bytes, bytes)

        df = pd.read_excel(io.BytesIO(excel_bytes), engine="openpyxl", sheet_name="Assessment Results")
        self.assertEqual(len(df), 2)
        self.assertEqual(df.iloc[0]["Skill"], "Ujian Gabungan")
        self.assertEqual(df.iloc[0]["Completed At"], "In Progress")
        self.assertEqual(df.iloc[1]["Skill"], "Skill 3")
        self.assertEqual(df.iloc[1]["Score"], "100.0%")
        self.assertEqual(df.iloc[1]["Duration"], "0m 17s")


if __name__ == "__main__":
    unittest.main()
