"""
Unit tests for Phase 12: DOCX Question Bank Import.
Tests:
- Programmatic generation of test Word (.docx) documents.
- Flexible parsing with python-docx across formatting variations.
- Detection of question text, options A-D, answer keys, explanations, and topics.
- Validation and error reporting for invalid questions (missing answer, < 4 options, duplicate options).
- Batch import execution saving questions as 'Draft' (never published automatically).
- Source file tracking on imported questions.
- Audit history logging in import_history table.
"""

import unittest
import tempfile
import io
from pathlib import Path
import docx

from database.db import init_db, get_connection
from database import queries
from utils.docx_parser import (
    parse_docx_content,
    extract_paragraphs_from_docx,
    infer_suggested_difficulty,
)
from services.import_service import (
    execute_import_batch,
    fetch_import_history_logs,
)


def create_sample_docx_bytes(include_invalid: bool = False) -> bytes:
    """Helper to generate an in-memory .docx file for testing."""
    doc = docx.Document()

    # Topic Header
    doc.add_paragraph("TOPIC: MODAL AUXILIARIES AND CONDITIONALS")

    # Question 1: Standard 'QUESTION 1' format
    doc.add_paragraph("QUESTION 1")
    doc.add_paragraph("If it rains tomorrow, we ___ our trip to the botanical garden.")
    doc.add_paragraph("A. cancel")
    doc.add_paragraph("B. will cancel")
    doc.add_paragraph("C. would cancel")
    doc.add_paragraph("D. had cancelled")
    doc.add_paragraph("ANSWER: B")
    doc.add_paragraph("EXPLANATION: First conditional requires 'will + base verb' in the main clause for a real future possibility.")

    # Question 2: '2.' format with inline options
    doc.add_paragraph("2. She would have passed the examination if she ___ more conscientiously.")
    doc.add_paragraph("A. studied   B. has studied   C. had studied   D. will study")
    doc.add_paragraph("Answer: C")
    doc.add_paragraph("Explanation: Third conditional requires past perfect 'had studied' in the if-clause.")

    # Question 3: 'Question 3:' with Indonesian-style 'Kunci:' and 'Pembahasan:'
    doc.add_paragraph("Question 3:")
    doc.add_paragraph("You ___ wear a helmet while riding a motorcycle according to city law.")
    doc.add_paragraph("A. might")
    doc.add_paragraph("B. must")
    doc.add_paragraph("C. could")
    doc.add_paragraph("D. may")
    doc.add_paragraph("Kunci: B")
    doc.add_paragraph("Pembahasan: 'Must' expresses legal obligation and requirement.")

    if include_invalid:
        # Invalid 1: Missing answer key
        doc.add_paragraph("QUESTION 4")
        doc.add_paragraph("The train ___ arrive on time despite the severe winter storm.")
        doc.add_paragraph("A. should")
        doc.add_paragraph("B. ought to")
        doc.add_paragraph("C. shall")
        doc.add_paragraph("D. will")
        # No answer line!
        doc.add_paragraph("EXPLANATION: Answer key was forgotten by the teacher.")

        # Invalid 2: Only 3 options (missing D)
        doc.add_paragraph("5. Water boils at 100 degrees Celsius under standard atmospheric pressure.")
        doc.add_paragraph("A. true statement")
        doc.add_paragraph("B. false statement")
        doc.add_paragraph("C. uncertain")
        doc.add_paragraph("ANSWER: A")
        doc.add_paragraph("EXPLANATION: Only 3 options provided.")

        # Invalid 3: Duplicate options
        doc.add_paragraph("QUESTION 6")
        doc.add_paragraph("Select the synonymous word for 'rapid':")
        doc.add_paragraph("A. fast")
        doc.add_paragraph("B. fast")  # Duplicate!
        doc.add_paragraph("C. slow")
        doc.add_paragraph("D. sluggish")
        doc.add_paragraph("ANSWER: A")
        doc.add_paragraph("EXPLANATION: Duplicate choices.")

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


class TestDocxQuestionBankImport(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp_dir.name) / "test_import.db"
        init_db(self.db_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    # ---------------------------------------------
    # 1. PARSER TESTS
    # ---------------------------------------------
    def test_parser_extracts_valid_docx_content(self):
        """Test parsing valid Word document with standard variations."""
        docx_bytes = create_sample_docx_bytes(include_invalid=False)
        result = parse_docx_content(docx_bytes)

        self.assertEqual(result["total_detected"], 3)
        self.assertEqual(result["valid_count"], 3)
        self.assertEqual(result["invalid_count"], 0)
        self.assertEqual(result["global_topic"], "MODAL AUXILIARIES AND CONDITIONALS")

        # Verify Question 1 parsed elements
        q1 = result["questions"][0]
        self.assertTrue(q1["is_valid"])
        self.assertIn("botanical garden", q1["question_text"])
        self.assertEqual(q1["option_a"], "cancel")
        self.assertEqual(q1["option_b"], "will cancel")
        self.assertEqual(q1["correct_answer"], "B")
        self.assertIn("First conditional", q1["explanation"])

        # Verify Question 2 with inline options
        q2 = result["questions"][1]
        self.assertTrue(q2["is_valid"])
        self.assertEqual(q2["correct_answer"], "C")
        self.assertTrue(bool(q2["option_a"]))
        self.assertTrue(bool(q2["option_b"]))

        # Verify Question 3 with Indonesian labels
        q3 = result["questions"][2]
        self.assertTrue(q3["is_valid"])
        self.assertEqual(q3["correct_answer"], "B")

    def test_parser_detects_and_reports_invalid_questions(self):
        """Test that invalid questions are flagged with explicit error descriptions."""
        docx_bytes = create_sample_docx_bytes(include_invalid=True)
        result = parse_docx_content(docx_bytes)

        self.assertEqual(result["total_detected"], 6)
        self.assertEqual(result["valid_count"], 3)
        self.assertEqual(result["invalid_count"], 3)

        invalid_list = [q for q in result["questions"] if not q["is_valid"]]
        self.assertEqual(len(invalid_list), 3)

        # Q4 missing answer key
        q4_errors = " ".join(invalid_list[0]["validation_errors"])
        self.assertIn("Answer key missing", q4_errors)

        # Q5 missing option D
        q5_errors = " ".join(invalid_list[1]["validation_errors"])
        self.assertIn("Missing option", q5_errors)

        # Q6 duplicate choices
        q6_errors = " ".join(invalid_list[2]["validation_errors"])
        self.assertIn("Duplicate choices", q6_errors)

    def test_infer_suggested_difficulty(self):
        """Test heuristic difficulty suggestion based on text complexity."""
        easy = infer_suggested_difficulty("Every day she goes to school.", "Simple present tense.")
        self.assertEqual(easy, "Easy")

        hard = infer_suggested_difficulty(
            "Nevertheless, hypothetical subjunctive inversion necessitates strict counter-illumination nuances.",
            "Complex historical grammar rule.",
        )
        self.assertEqual(hard, "Hard")

    # ---------------------------------------------
    # 2. BATCH IMPORT EXECUTION TESTS
    # ---------------------------------------------
    def test_batch_import_saves_drafts_and_never_publishes_automatically(self):
        """Verify that imported questions default to 'Draft' and track source_file."""
        docx_bytes = create_sample_docx_bytes(include_invalid=True)
        result = parse_docx_content(docx_bytes)

        detected, imported, failed = execute_import_batch(
            file_name="Modul_Grammar_Conditionals_Test.docx",
            questions=result["questions"],
            category_name="Grammar",
            default_status="Draft",
            imported_by="test_admin",
            db_path=self.db_path,
        )

        self.assertEqual(detected, 6)
        self.assertEqual(imported, 3)  # Only 3 valid questions imported
        self.assertEqual(failed, 3)    # 3 invalid questions skipped

        # Inspect database
        with get_connection(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT question_text, status, source_file FROM questions WHERE source_file = ?;", ("Modul_Grammar_Conditionals_Test.docx",))
            rows = cursor.fetchall()
            self.assertEqual(len(rows), 3)

            for r in rows:
                self.assertEqual(r[1], "Draft", "Imported questions must default to 'Draft' status.")
                self.assertEqual(r[2], "Modul_Grammar_Conditionals_Test.docx")

    # ---------------------------------------------
    # 3. IMPORT HISTORY AUDIT TESTS
    # ---------------------------------------------
    def test_import_history_audit_logging(self):
        """Test that import batch executions are recorded in import_history table."""
        docx_bytes = create_sample_docx_bytes(include_invalid=False)
        result = parse_docx_content(docx_bytes)

        file_name = "Import_Audit_Test_File.docx"
        execute_import_batch(
            file_name=file_name,
            questions=result["questions"],
            category_name="Grammar",
            imported_by="superadmin_test",
            db_path=self.db_path,
        )

        logs = fetch_import_history_logs(limit=10, db_path=self.db_path)
        self.assertGreaterEqual(len(logs), 1)

        first_log = logs[0]
        self.assertEqual(first_log.file_name, file_name)
        self.assertEqual(first_log.category, "Grammar")
        self.assertEqual(first_log.total_detected, 3)
        self.assertEqual(first_log.total_imported, 3)
        self.assertEqual(first_log.total_failed, 0)
        self.assertEqual(first_log.imported_by, "superadmin_test")
        self.assertIsNotNone(first_log.imported_at)


if __name__ == "__main__":
    unittest.main()
